# -*- coding: utf-8 -*-
"""Text edition (.docx + .txt): every card with clickable links, for Facebook captions."""
import json, re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
J=json.load(open('/home/claude/work/build/nuqul_7_9.json',encoding='utf8'))
SN=J['meta']['sections']; ORDER=J['meta']['section_order']
cards=sorted([e for e in J['entries'] if e.get('card_number')],key=lambda e:e['card_number'])
SUP=J['meta']['legend']['supports']
FONT='Calibri'
def is_ar(s): return bool(re.search('[؀-ۿ]',s or ''))
def set_rtl(p,rtl=True):
    pPr=p._p.get_or_add_pPr()
    b=OxmlElement('w:bidi'); b.set(qn('w:val'),'1' if rtl else '0'); pPr.append(b)
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
def run(p,text,size=14,bold=False,color=None,rtl=True):
    r=p.add_run(text); r.font.name=FONT; r.font.size=Pt(size); r.bold=bold
    rPr=r._r.get_or_add_rPr()
    f=rPr.find(qn('w:rFonts'))
    if f is None: f=OxmlElement('w:rFonts'); rPr.append(f)
    for a in ('w:ascii','w:hAnsi','w:cs'): f.set(qn(a),FONT)
    sz=OxmlElement('w:szCs'); sz.set(qn('w:val'),str(int(size*2))); rPr.append(sz)
    if bold: bc=OxmlElement('w:bCs'); rPr.append(bc)
    if rtl: rt=OxmlElement('w:rtl'); rPr.append(rt)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    return r
def para(doc,rtl=True,space_after=4,space_before=0):
    p=doc.add_paragraph(); set_rtl(p,rtl)
    pf=p.paragraph_format; pf.space_after=Pt(space_after); pf.space_before=Pt(space_before); pf.line_spacing=1.4
    return p
def hyperlink(p,url,text,size=13):
    rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid)
    r=OxmlElement('w:r'); rPr=OxmlElement('w:rPr')
    f=OxmlElement('w:rFonts')
    for a in ('w:ascii','w:hAnsi','w:cs'): f.set(qn(a),FONT)
    rPr.append(f)
    c=OxmlElement('w:color'); c.set(qn('w:val'),'1F5FA8'); rPr.append(c)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rPr.append(u)
    s=OxmlElement('w:sz'); s.set(qn('w:val'),str(size*2)); rPr.append(s)
    s2=OxmlElement('w:szCs'); s2.set(qn('w:val'),str(size*2)); rPr.append(s2)
    r.append(rPr); t=OxmlElement('w:t'); t.text=text; t.set(qn('xml:space'),'preserve'); r.append(t); h.append(r)
    p._p.append(h)
def label_line(doc,label,text,color='7E000B',size=14):
    if not text: return None
    p=para(doc); run(p,label+': ',size,True,color); run(p,text,size,rtl=True); return p

doc=Document()
sec=doc.sections[0]; sec.left_margin=sec.right_margin=Cm(2); sec.top_margin=sec.bottom_margin=Cm(2)
st=doc.styles['Normal']; st.font.name=FONT; st.font.size=Pt(14)
p=para(doc,space_after=6); run(p,'حول سن الحضانة والرعاية المشتركة بالمناصفة — النصوص والمصادر',22,True,'7E000B')
p=para(doc,space_after=4); run(p,f'الطبعة النصية لكل البطاقات بروابطها القابلة للنقر · الإصدار {J["meta"]["version"]} · {len(cards)} بطاقة في {len(ORDER)} قسما',14,False,'4A5A61')
p=para(doc,space_after=4); run(p,'ما بين معقوفين [ ] إيضاح لمرجع الضمير، وليس من لفظ المصدر. لكل بطاقة نصها الأصلي وترجمته ومصدره ومنزلته ووجه الاستدلال به وروابط التحقق منه.',13,False,'4A5A61')
TXT=['حول سن الحضانة والرعاية المشتركة بالمناصفة — النصوص والمصادر',f'الطبعة النصية · الإصدار {J["meta"]["version"]} · {len(cards)} بطاقة','']
cur=None
SEP='━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━'
for e in cards:
    if e['section']!=cur:
        cur=e['section']
        p=para(doc,space_before=30,space_after=10); p.paragraph_format.page_break_before=True
        run(p,f'القسم {ORDER.index(cur)+1:02d}: {SN[cur]}',20,True,'1F5FA8')
        TXT+=['','',f'■■■ القسم {ORDER.index(cur)+1:02d}: {SN[cur]} ■■■','']
    n=e['card_number']
    p=para(doc,space_before=6,space_after=6); run(p,f'بطاقة {n}',16,True,'C40000'); run(p,f' · {SN[cur]}',12,False,'4A5A61')
    p=para(doc,space_after=8); run(p,e['title'],17,True,'182228')
    o=e['original']; ot=o['text']; ar=is_ar(ot[:60])
    shown=ot if (not ar or ot.startswith('مضمونه')) else ('«'+ot.strip('«»')+'»')
    p=para(doc,rtl=ar,space_after=6); run(p,shown,15,False,'182228',rtl=ar)
    tr=e.get('translation_ar') or ''
    if tr: label_line(doc,'الترجمة',tr)
    a=e['author']; w=e['work']
    src=' · '.join(x for x in [a['name']+(f' ({a["date"]})' if a.get('date') else ''), w.get('title'), w.get('section'), w.get('volume'), w.get('page'), w.get('edition'), e.get('jurisdiction')] if x)
    label_line(doc,'المصدر',src,'1F5FA8',13)
    about=e.get('source_standing') or e.get('about') or ''
    if about: label_line(doc,'عن المصدر',about+(f' — درجة الدليل: {e["evidence_level"]}' if e.get('evidence_level') else ''),'8A6418',13)
    label_line(doc,'الاستدلال',e['argument'],'0F5C3A',14)
    links=e.get('links') or [{'url':u,'label':'الرابط'} for u in e.get('urls',[])]
    tl=[]
    for l in links:
        p=para(doc,space_after=2); run(p,'🔗 '+(l.get('label') or 'الرابط')+': ',13,True,'4A5A61'); hyperlink(p,l['url'],l['url'])
        tl.append(f"🔗 {l.get('label') or 'الرابط'}: {l['url']}")
    sup=' '.join('#'+SUP[s].replace(' ','_') for s in e['supports'] if s in SUP)
    p=para(doc,space_after=2); run(p,sup+' #سن_الحضانة #حق_الطفل_في_أبويه',12,False,'5B4B9A')
    # wide separator
    for k in range(2): para(doc,space_after=6)
    p=para(doc,space_after=6); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,SEP,14,False,'B9B0A0')
    for k in range(2): para(doc,space_after=6)
    TXT+=[f'بطاقة {n} · {SN[cur]}',e['title'],'',shown]
    if tr: TXT+=['',f'الترجمة: {tr}']
    TXT+=['',f'المصدر: {src}']
    if about: TXT+=[f'عن المصدر: {about}'+(f' — درجة الدليل: {e["evidence_level"]}' if e.get('evidence_level') else '')]
    TXT+=['',f'الاستدلال: {e["argument"]}','']+tl+[sup+' #سن_الحضانة #حق_الطفل_في_أبويه','','',SEP,'','','']
out='/home/claude/work/build/الطبعة_النصية_للبطاقات_بالروابط'
doc.save(out+'.docx'); open(out+'.txt','w',encoding='utf-8-sig').write('\n'.join(TXT))
print(len(cards))
